#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script sinh file Báo cáo Nhân lực hàng ngày từ dữ liệu JSON.
Sử dụng:
    python3 generate_report.py [tuy_chon_file_json] [tuy_chon_file_xuat.xlsx]
Mặc định:
    python3 generate_report.py data_mau.json Bao_Cao_Nhan_Luc_Hang_Ngay.xlsx
"""

import sys
import os
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image

def generate_excel(json_path="data_mau.json", output_path="Bao_Cao_Nhan_Luc_Hang_Ngay.xlsx"):
    if not os.path.exists(json_path):
        print(f"Lỗi: Không tìm thấy file dữ liệu {json_path}")
        return

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo Cáo Nhân Lực"

    # Color definitions
    NAVY = "0B3064"
    NAVY_LIGHT = "EBF1F8"
    RED = "B92521"
    BORDER_COLOR = "9FB8D4"
    WHITE = "FFFFFF"

    font_title = Font(name="Arial", size=15, bold=True, color=NAVY)
    font_sub_title = Font(name="Arial", size=11, bold=True, color=NAVY)
    font_meta_label = Font(name="Arial", size=9.5, bold=True, color=NAVY)
    font_meta_val = Font(name="Arial", size=9.5, bold=True, color=RED)
    font_sec_hdr = Font(name="Arial", size=10, bold=True, color=WHITE)
    font_tbl_hdr = Font(name="Arial", size=9, bold=True, color=NAVY)
    font_tbl_data = Font(name="Arial", size=9.5, color="000000")
    font_tbl_num = Font(name="Arial", size=10, bold=True, color=RED)
    font_total_lbl = Font(name="Arial", size=10, bold=True, color=NAVY)
    font_total_num = Font(name="Arial", size=12, bold=True, color=RED)

    fill_sec_hdr = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    fill_tbl_hdr = PatternFill(start_color=NAVY_LIGHT, end_color=NAVY_LIGHT, fill_type="solid")
    fill_total = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border_side = Side(style="thin", color=BORDER_COLOR)
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Column widths
    ws.column_dimensions['A'].width = 5.5
    ws.column_dimensions['B'].width = 17
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 17
    ws.column_dimensions['E'].width = 9
    ws.column_dimensions['F'].width = 2.5
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 33
    ws.column_dimensions['I'].width = 15

    # Header Left
    ws.merge_cells('A1:E1')
    ws['A1'] = "BÁO CÁO NHÂN LỰC HÀNG NGÀY"
    ws['A1'].font = font_title
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A2:E2')
    ws['A2'] = "TẠI CÔNG TRƯỜNG"
    ws['A2'].font = font_sub_title
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")

    ws['A3'] = "Dự án:"
    ws['A3'].font = font_meta_label
    ws.merge_cells('B3:E3')
    ws['B3'] = data.get("du_an", "")
    ws['B3'].font = font_meta_val

    ws['A4'] = "Gói thầu:"
    ws['A4'].font = font_meta_label
    ws.merge_cells('B4:E4')
    ws['B4'] = data.get("goi_thau", "")
    ws['B4'].font = font_meta_val

    ws['A5'] = "Nhà thầu:"
    ws['A5'].font = font_meta_label
    ws.merge_cells('B5:C5')
    ws['B5'] = data.get("nha_thau", "")
    ws['B5'].font = font_meta_val

    ws['D5'] = "Ngày báo cáo:"
    ws['D5'].font = font_meta_label
    ws['D5'].alignment = Alignment(horizontal="right")
    ws['E5'] = data.get("ngay_bao_cao", "")
    ws['E5'].font = font_meta_val
    ws['E5'].alignment = Alignment(horizontal="center")

    wt = data.get("thoi_tiet", {})
    wt_str = f" [{'x' if wt.get('nang') else ' '}] Nắng    [{'x' if wt.get('mua') else ' '}] Mưa    [{'x' if wt.get('am_u') else ' '}] Âm u"
    if wt.get("khac"):
        wt_str += f"    Khác: {wt.get('khac')}"
    else:
        wt_str += "    Khác: ................."
    ws['A6'] = "Thời tiết:"
    ws['A6'].font = font_meta_label
    ws.merge_cells('B6:E6')
    ws['B6'] = wt_str
    ws['B6'].font = Font(name="Arial", size=9, bold=False, color="1C2430")

    # Header Right (Banner)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    banner_path = os.path.join(script_dir, "assets", "banner_cienco510.png")
    if os.path.exists(banner_path):
        img = Image(banner_path)
        img.width = 390
        img.height = 110
        ws.add_image(img, 'G1')

    # Separator
    for col in range(1, 10):
        ws.cell(row=7, column=col).border = Border(bottom=thin_border_side)

    # 1. THỐNG KÊ NHÂN LỰC
    ws.merge_cells('A8:E8')
    ws['A8'] = "1. THỐNG KÊ NHÂN LỰC"
    ws['A8'].font = font_sec_hdr
    ws['A8'].fill = fill_sec_hdr
    ws['A8'].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws['A9'] = "STT"
    ws['A9'].font = font_tbl_hdr
    ws['A9'].fill = fill_tbl_hdr
    ws['A9'].alignment = Alignment(horizontal="center", vertical="center")
    ws['A9'].border = box_border

    ws.merge_cells('B9:D9')
    ws['B9'] = "BỘ PHẬN / TỔ ĐỘI"
    ws['B9'].font = font_tbl_hdr
    ws['B9'].fill = fill_tbl_hdr
    ws['B9'].alignment = Alignment(horizontal="center", vertical="center")
    for col in ['B', 'C', 'D']:
        ws[f'{col}9'].border = box_border

    ws['E9'] = "SỐ LƯỢNG (NGƯỜI)"
    ws['E9'].font = font_tbl_hdr
    ws['E9'].fill = fill_tbl_hdr
    ws['E9'].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws['E9'].border = box_border

    sec1_list = data.get("thong_ke_nhan_luc", [])
    row_start_sec1 = 10
    for idx, item in enumerate(sec1_list, start=1):
        r = row_start_sec1 + idx - 1
        ws[f'A{r}'] = idx
        ws[f'A{r}'].alignment = Alignment(horizontal="center", vertical="center")
        ws[f'A{r}'].font = font_tbl_data
        ws[f'A{r}'].border = box_border

        ws.merge_cells(f'B{r}:D{r}')
        ws[f'B{r}'] = item.get("bo_phan", "")
        ws[f'B{r}'].font = font_tbl_data
        ws[f'B{r}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in ['B', 'C', 'D']:
            ws[f'{col}{r}'].border = box_border

        ws[f'E{r}'] = item.get("so_luong", 0)
        ws[f'E{r}'].font = font_tbl_num
        ws[f'E{r}'].alignment = Alignment(horizontal="center", vertical="center")
        ws[f'E{r}'].border = box_border

    row_end_sec1 = row_start_sec1 + len(sec1_list) - 1
    total_sec1_row = row_end_sec1 + 1

    ws.merge_cells(f'A{total_sec1_row}:D{total_sec1_row}')
    ws[f'A{total_sec1_row}'] = "TỔNG CỘNG"
    ws[f'A{total_sec1_row}'].font = font_total_lbl
    ws[f'A{total_sec1_row}'].fill = fill_total
    ws[f'A{total_sec1_row}'].alignment = Alignment(horizontal="center", vertical="center")
    for col in ['A', 'B', 'C', 'D']:
        ws[f'{col}{total_sec1_row}'].border = box_border

    ws[f'E{total_sec1_row}'] = f"=SUM(E{row_start_sec1}:E{row_end_sec1})"
    ws[f'E{total_sec1_row}'].font = font_total_num
    ws[f'E{total_sec1_row}'].fill = fill_total
    ws[f'E{total_sec1_row}'].alignment = Alignment(horizontal="center", vertical="center")
    ws[f'E{total_sec1_row}'].border = box_border

    # 4. KẾ HOẠCH NHÂN LỰC NGÀY MAI
    sec4_header_row = total_sec1_row + 2
    kh = data.get("ke_hoach_ngay_mai", {})
    ngay_mai = kh.get("ngay", "")
    ws.merge_cells(f'A{sec4_header_row}:E{sec4_header_row}')
    ws[f'A{sec4_header_row}'] = f"4. KẾ HOẠCH NHÂN LỰC NGÀY MAI ({ngay_mai})" if ngay_mai else "4. KẾ HOẠCH NHÂN LỰC NGÀY MAI"
    ws[f'A{sec4_header_row}'].font = font_sec_hdr
    ws[f'A{sec4_header_row}'].fill = fill_sec_hdr
    ws[f'A{sec4_header_row}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    sec4_tbl_hdr = sec4_header_row + 1
    ws.merge_cells(f'A{sec4_tbl_hdr}:B{sec4_tbl_hdr}')
    ws[f'A{sec4_tbl_hdr}'] = "BỘ PHẬN / TỔ ĐỘI"
    ws[f'A{sec4_tbl_hdr}'].font = font_tbl_hdr
    ws[f'A{sec4_tbl_hdr}'].fill = fill_tbl_hdr
    ws[f'A{sec4_tbl_hdr}'].alignment = Alignment(horizontal="center", vertical="center")
    ws[f'A{sec4_tbl_hdr}'].border = box_border
    ws[f'B{sec4_tbl_hdr}'].border = box_border

    ws[f'C{sec4_tbl_hdr}'] = "NHÂN LỰC DỰ KIẾN"
    ws[f'C{sec4_tbl_hdr}'].font = font_tbl_hdr
    ws[f'C{sec4_tbl_hdr}'].fill = fill_tbl_hdr
    ws[f'C{sec4_tbl_hdr}'].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws[f'C{sec4_tbl_hdr}'].border = box_border

    ws[f'D{sec4_tbl_hdr}'] = "BỘ PHẬN / TỔ ĐỘI"
    ws[f'D{sec4_tbl_hdr}'].font = font_tbl_hdr
    ws[f'D{sec4_tbl_hdr}'].fill = fill_tbl_hdr
    ws[f'D{sec4_tbl_hdr}'].alignment = Alignment(horizontal="center", vertical="center")
    ws[f'D{sec4_tbl_hdr}'].border = box_border

    ws[f'E{sec4_tbl_hdr}'] = "NHÂN LỰC DỰ KIẾN"
    ws[f'E{sec4_tbl_hdr}'].font = font_tbl_hdr
    ws[f'E{sec4_tbl_hdr}'].fill = fill_tbl_hdr
    ws[f'E{sec4_tbl_hdr}'].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws[f'E{sec4_tbl_hdr}'].border = box_border

    trai = kh.get("danh_sach_trai", [])
    phai = kh.get("danh_sach_phai", [])
    max_len = max(len(trai), len(phai), 1)

    sec4_start_row = sec4_tbl_hdr + 1
    for i in range(max_len):
        r = sec4_start_row + i
        it_l = trai[i] if i < len(trai) else {}
        it_r = phai[i] if i < len(phai) else {}

        ws.merge_cells(f'A{r}:B{r}')
        ws[f'A{r}'] = it_l.get("bo_phan", "")
        ws[f'A{r}'].font = font_tbl_data
        ws[f'A{r}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws[f'A{r}'].border = box_border
        ws[f'B{r}'].border = box_border

        ws[f'C{r}'] = it_l.get("so_luong", "")
        ws[f'C{r}'].font = font_tbl_num
        ws[f'C{r}'].alignment = Alignment(horizontal="center", vertical="center")
        ws[f'C{r}'].border = box_border

        ws[f'D{r}'] = it_r.get("bo_phan", "")
        ws[f'D{r}'].font = font_tbl_data
        ws[f'D{r}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws[f'D{r}'].border = box_border

        ws[f'E{r}'] = it_r.get("so_luong", "")
        ws[f'E{r}'].font = font_tbl_num
        ws[f'E{r}'].alignment = Alignment(horizontal="center", vertical="center")
        ws[f'E{r}'].border = box_border

    sec4_end_row = sec4_start_row + max_len - 1
    sec4_total_row = sec4_end_row + 1

    ws.merge_cells(f'A{sec4_total_row}:D{sec4_total_row}')
    ws[f'A{sec4_total_row}'] = "TỔNG CỘNG"
    ws[f'A{sec4_total_row}'].font = font_total_num
    ws[f'A{sec4_total_row}'].fill = fill_total
    ws[f'A{sec4_total_row}'].alignment = Alignment(horizontal="right", vertical="center")
    for col in ['A', 'B', 'C', 'D']:
        ws[f'{col}{sec4_total_row}'].border = box_border

    ws[f'E{sec4_total_row}'] = f"=SUM(C{sec4_start_row}:C{sec4_end_row})+SUM(E{sec4_start_row}:E{sec4_end_row})"
    ws[f'E{sec4_total_row}'].font = font_total_num
    ws[f'E{sec4_total_row}'].fill = fill_total
    ws[f'E{sec4_total_row}'].alignment = Alignment(horizontal="center", vertical="center")
    ws[f'E{sec4_total_row}'].border = box_border

    # 2. NHÂN LỰC THEO KHU VỰC THI CÔNG
    ws.merge_cells('G8:I8')
    ws['G8'] = "2. NHÂN LỰC THEO KHU VỰC THI CÔNG"
    ws['G8'].font = font_sec_hdr
    ws['G8'].fill = fill_sec_hdr
    ws['G8'].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws['G9'] = "KHU VỰC"
    ws['G9'].font = font_tbl_hdr
    ws['G9'].fill = fill_tbl_hdr
    ws['G9'].alignment = Alignment(horizontal="center", vertical="center")
    ws['G9'].border = box_border

    ws['H9'] = "NỘI DUNG CÔNG VIỆC"
    ws['H9'].font = font_tbl_hdr
    ws['H9'].fill = fill_tbl_hdr
    ws['H9'].alignment = Alignment(horizontal="center", vertical="center")
    ws['H9'].border = box_border

    ws['I9'] = "SỐ LƯỢNG NHÂN CÔNG"
    ws['I9'].font = font_tbl_hdr
    ws['I9'].fill = fill_tbl_hdr
    ws['I9'].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws['I9'].border = box_border

    sec2_list = data.get("nhan_luc_theo_khu_vuc", [])
    for idx, item in enumerate(sec2_list, start=1):
        r = 9 + idx
        ws[f'G{r}'] = item.get("khu_vuc", "")
        ws[f'G{r}'].alignment = Alignment(horizontal="center", vertical="center")
        ws[f'G{r}'].font = font_tbl_data
        ws[f'G{r}'].border = box_border

        ws[f'H{r}'] = item.get("noi_dung", "")
        ws[f'H{r}'].font = font_tbl_data
        ws[f'H{r}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws[f'H{r}'].border = box_border

        ws[f'I{r}'] = item.get("so_luong", 0)
        ws[f'I{r}'].font = font_tbl_num
        ws[f'I{r}'].alignment = Alignment(horizontal="center", vertical="center")
        ws[f'I{r}'].border = box_border

    # 3. ĐÁNH GIÁ TÌNH HÌNH NHÂN LỰC
    eval_hdr_row = 10 + len(sec2_list)
    ws.merge_cells(f'G{eval_hdr_row}:I{eval_hdr_row}')
    ws[f'G{eval_hdr_row}'] = "3. ĐÁNH GIÁ TÌNH HÌNH NHÂN LỰC"
    ws[f'G{eval_hdr_row}'].font = font_sec_hdr
    ws[f'G{eval_hdr_row}'].fill = fill_sec_hdr
    ws[f'G{eval_hdr_row}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ev = data.get("danh_gia", {})
    tien_do_str = " [x] Có     [ ] Không" if ev.get("dap_ung_tien_do") == "co" else " [ ] Có     [x] Không"
    atld_val = ev.get("an_toan_lao_dong", "tot")
    atld_str = f" [{'x' if atld_val=='tot' else ' '}] Tốt   [{'x' if atld_val=='kha' else ' '}] Khá   [{'x' if atld_val=='tb' else ' '}] Trung bình   [{'x' if atld_val=='kem' else ' '}] Kém"

    eval_rows = [
        ('• Tổng số nhân lực hiện trường:', f'=E{total_sec1_row} & " người."'),
        ('• Nhân lực đáp ứng tiến độ:', tien_do_str),
        ('• Tình hình chấp hành ATLĐ:', atld_str),
        ('• Số nhân công mới vào công trường:', f' {ev.get("nhan_cong_moi", "00")} người.'),
        ('• Số nhân công nghỉ việc / nghỉ phép:', f' {ev.get("nhan_cong_nghi", "00")} người.'),
        ('• Ghi chú khác:', f' {ev.get("ghi_chu_khac") if ev.get("ghi_chu_khac") else "............................................................................"}')
    ]

    r_cur = eval_hdr_row + 1
    eval_start = r_cur
    for label, val in eval_rows:
        ws[f'G{r_cur}'] = label
        ws[f'G{r_cur}'].font = Font(name="Arial", size=9, bold=False, color="1C2430")
        ws[f'G{r_cur}'].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells(f'H{r_cur}:I{r_cur}')
        if val.startswith('='):
            ws[f'H{r_cur}'] = val
            ws[f'H{r_cur}'].font = font_tbl_num
        else:
            ws[f'H{r_cur}'] = val
            ws[f'H{r_cur}'].font = Font(name="Arial", size=9, bold=False, color="1C2430")
        ws[f'H{r_cur}'].alignment = Alignment(horizontal="left", vertical="center")

        ws[f'G{r_cur}'].border = Border(left=thin_border_side)
        ws[f'I{r_cur}'].border = Border(right=thin_border_side)
        r_cur += 1

    eval_end = r_cur - 1
    for c in ['G', 'H', 'I']:
        ws[f'{c}{eval_start}'].border = Border(top=thin_border_side, left=ws[f'{c}{eval_start}'].border.left, right=ws[f'{c}{eval_start}'].border.right)
        ws[f'{c}{eval_end}'].border = Border(bottom=thin_border_side, left=ws[f'{c}{eval_end}'].border.left, right=ws[f'{c}{eval_end}'].border.right)

    # 5. XÁC NHẬN
    sig_hdr_row = eval_end + 1
    ws.merge_cells(f'G{sig_hdr_row}:I{sig_hdr_row}')
    ws[f'G{sig_hdr_row}'] = "XÁC NHẬN"
    ws[f'G{sig_hdr_row}'].font = font_sec_hdr
    ws[f'G{sig_hdr_row}'].fill = fill_sec_hdr
    ws[f'G{sig_hdr_row}'].alignment = Alignment(horizontal="center", vertical="center")

    sig_title_row = sig_hdr_row + 1
    xn = data.get("xac_nhan", {})
    ws.merge_cells(f'G{sig_title_row}:I{sig_title_row}')
    ws[f'G{sig_title_row}'] = xn.get("chuc_vu", "CHỈ HUY TRƯỞNG CÔNG TRƯỜNG")
    ws[f'G{sig_title_row}'].font = font_sub_title
    ws[f'G{sig_title_row}'].alignment = Alignment(horizontal="center", vertical="center")

    sig_img_row = sig_title_row + 1
    ws.row_dimensions[sig_img_row].height = 36
    ws.merge_cells(f'G{sig_img_row}:I{sig_img_row}')

    sig_path = os.path.join(script_dir, xn.get("chu_ky_path", "assets/signature_sample.png"))
    if os.path.exists(sig_path):
        sig_img = Image(sig_path)
        sig_img.width = 120
        sig_img.height = 42
        ws.add_image(sig_img, f'H{sig_img_row}')

    sig_name_row = sig_img_row + 1
    ws.merge_cells(f'G{sig_name_row}:I{sig_name_row}')
    ws[f'G{sig_name_row}'] = xn.get("ho_ten", "Nguyễn Anh Tuấn")
    ws[f'G{sig_name_row}'].font = font_sub_title
    ws[f'G{sig_name_row}'].alignment = Alignment(horizontal="center", vertical="center")

    for r in range(sig_title_row, sig_name_row + 1):
        ws[f'G{r}'].border = Border(left=thin_border_side)
        ws[f'I{r}'].border = Border(right=thin_border_side)
    for c in ['G', 'H', 'I']:
        ws[f'{c}{sig_name_row}'].border = Border(bottom=thin_border_side, left=ws[f'{c}{sig_name_row}'].border.left, right=ws[f'{c}{sig_name_row}'].border.right)

    # Page setup
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.page_setup.horizontalCentered = True
    ws.views.sheetView[0].showGridLines = True

    wb.save(output_path)
    print(f"Đã xuất thành công báo cáo Excel: {output_path}")

if __name__ == "__main__":
    json_f = sys.argv[1] if len(sys.argv) > 1 else "data_mau.json"
    xlsx_f = sys.argv[2] if len(sys.argv) > 2 else "Bao_Cao_Nhan_Luc_Hang_Ngay.xlsx"
    generate_excel(json_f, xlsx_f)
